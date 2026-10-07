"""All PDF operations. Each function takes input files + parameters
and returns a result (bytes + file name) or JSON data."""
import base64
import difflib
import io
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import datetime

import pymupdf
from PIL import Image

from scan import scanify_pdf


class ToolError(Exception):
    pass


MM = 72 / 25.4
PAPER = {"a4": (595.28, 841.89), "letter": (612, 792), "a5": (419.53, 595.28), "a3": (841.89, 1190.55)}
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff", ".webp", ".heic"}
OFFICE_EXT = {".doc", ".docx", ".odt", ".rtf", ".txt", ".xls", ".xlsx", ".ods", ".csv",
              ".ppt", ".pptx", ".odp", ".html", ".htm"}
FONTS = {"helv", "hebo", "heit", "hebi", "tiro", "tibo", "tiit", "tibi", "cour", "cobo", "coit", "cobi"}


# ---------------------------------------------------------------- Helfer

NAME_EN = {"_zusammengefuegt": "_merged", "_geteilt": "_split", "_bereinigt": "_pages-removed", "_auszug": "_extract",
           "_seiten": "_pages", "_organisiert": "_organized", "_gedreht": "_rotated", "_komprimiert": "_compressed",
           "_unterschrieben_scan": "_signed_scanned", "_scan": "_scanned", "_graustufen": "_grayscale", "_bilder": "_images",
           "_nummeriert": "_numbered", "_wasserzeichen": "_watermarked", "_zugeschnitten": "_cropped",
           "_geschuetzt": "_protected", "_entsperrt": "_unlocked", "_geschwaerzt": "_redacted", "_repariert": "_repaired",
           "_flach": "_flattened", "_ausgefuellt": "_filled", "_bearbeitet": "_edited", "konvertiert": "converted"}


def localize_name(name, lang):
    """Result file names for the English UI: A_zusammengefuegt.pdf → A_merged.pdf"""
    if lang != "en":
        return name
    base, ext = os.path.splitext(name)
    for de, en in NAME_EN.items():
        if base.endswith(de):
            return base[: -len(de)] + en + ext
    return name


def stem(name):
    return os.path.splitext(os.path.basename(name))[0]


def open_pdf(f, password=None):
    try:
        doc = pymupdf.open(f["path"])
    except Exception as e:
        raise ToolError(f"„{f['name']}“ ist kein lesbares PDF ({e}).")
    if doc.needs_pass:
        if not password or not doc.authenticate(password):
            raise ToolError(f"„{f['name']}“ ist passwortgeschützt. Bitte zuerst mit „PDF entsperren“ öffnen.")
    return doc


def parse_pages(spec, n, default_all=True):
    """'1-3, 5, 8-' -> [0,1,2,4,7,...]. 1-basiert, 'ende'/'last' erlaubt."""
    spec = (spec or "").strip().lower()
    if not spec:
        if default_all:
            return list(range(n))
        raise ToolError("Bitte Seiten angeben (z. B. 1-3, 5).")
    spec = spec.replace("ende", str(n)).replace("last", str(n)).replace("end", str(n)).replace("–", "-")
    out = []
    for part in re.split(r"[,;\s]+", spec):
        if not part:
            continue
        m = re.fullmatch(r"(\d*)-(\d*)", part)
        if m:
            a = int(m.group(1)) if m.group(1) else 1
            b = int(m.group(2)) if m.group(2) else n
        elif part.isdigit():
            a = b = int(part)
        else:
            raise ToolError(f"Ungültige Seitenangabe: „{part}“")
        if a > b:
            a, b = b, a
        for p in range(a, b + 1):
            if 1 <= p <= n and (p - 1) not in out:
                out.append(p - 1)
    if not out:
        raise ToolError(f"Keine gültigen Seiten (Dokument hat {n} Seiten).")
    return out


def parse_groups(spec, n):
    # „1-3, 4-6, 7“ -> drei Gruppen; „1,3,5“ -> drei Einzelseiten
    groups = [g for part in (spec or "").split(",") for g in [part.strip()] if g]
    return [parse_pages(g, n, default_all=False) for g in groups]


def save_bytes(doc, **kw):
    opts = dict(garbage=3, deflate=True)
    opts.update(kw)
    return doc.tobytes(**opts)


def hex_to_rgb(h, default=(0, 0, 0)):
    h = (h or "").lstrip("#")
    if len(h) != 6:
        return default
    return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))


def zip_files(files):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files:
            z.writestr(name, data)
    return buf.getvalue()


def pdf_result(doc, name, info=None, **savekw):
    return {"data": save_bytes(doc, **savekw), "name": name, "info": info}


def image_to_pdf_page(out, path, size="fit", orientation="auto", margin=0):
    img = pymupdf.open(path)
    if img.is_pdf:
        out.insert_pdf(img)
        return
    rect = img[0].rect
    if size == "fit":
        w, h = rect.width, rect.height
    else:
        w, h = PAPER.get(size, PAPER["a4"])
        if orientation == "landscape" or (orientation == "auto" and rect.width > rect.height):
            w, h = h, w
    page = out.new_page(width=w, height=h)
    m = margin if size != "fit" else 0
    target = pymupdf.Rect(m, m, w - m, h - m)
    pdfbytes = img.convert_to_pdf()
    src = pymupdf.open("pdf", pdfbytes)
    page.show_pdf_page(target, src, 0, keep_proportion=True)
    img.close()


def load_image_as_pdf_path(path):
    """Convert HEIC & co. to PNG via Pillow if MuPDF cannot read them."""
    try:
        d = pymupdf.open(path)
        d.close()
        return path
    except Exception:
        img = Image.open(path)
        tmp = path + ".png"
        img.convert("RGB").save(tmp)
        return tmp


# ---------------------------------------------------------------- Werkzeuge

def merge(files, p):
    if len(files) < 2:
        raise ToolError("Bitte mindestens zwei Dateien auswählen.")
    out = pymupdf.open()
    for f in files:
        ext = os.path.splitext(f["name"])[1].lower()
        if ext in IMAGE_EXT:
            image_to_pdf_page(out, load_image_as_pdf_path(f["path"]), size="a4", margin=0)
            continue
        d = open_pdf(f)
        out.insert_pdf(d)
        d.close()
    if p.get("bookmarks"):
        toc, page = [], 1
        for f in files:
            ext = os.path.splitext(f["name"])[1].lower()
            n = 1 if ext in IMAGE_EXT else pymupdf.open(f["path"]).page_count
            toc.append([1, stem(f["name"]), page])
            page += n
        out.set_toc(toc)
    return pdf_result(out, f"{stem(files[0]['name'])}_zusammengefuegt.pdf",
                      f"{len(files)} Dateien → {out.page_count} Seiten")


def split(files, p):
    f = files[0]
    doc = open_pdf(f)
    n = doc.page_count
    mode = p.get("mode", "ranges")
    if mode == "every":
        k = max(1, int(p.get("every") or 1))
        groups = [list(range(i, min(i + k, n))) for i in range(0, n, k)]
    elif mode == "single":
        groups = [[i] for i in range(n)]
    else:
        groups = parse_groups(p.get("ranges"), n)
    parts = []
    for g in groups:
        part = pymupdf.open()
        for i in g:
            part.insert_pdf(doc, from_page=i, to_page=i)
        label = f"{g[0] + 1}" if len(g) == 1 else f"{g[0] + 1}-{g[-1] + 1}"
        parts.append((f"{stem(f['name'])}_S{label}.pdf", save_bytes(part)))
    if len(parts) == 1:
        return {"data": parts[0][1], "name": parts[0][0], "info": "1 Datei erstellt"}
    return {"data": zip_files(parts), "name": f"{stem(f['name'])}_geteilt.zip",
            "info": f"{len(parts)} PDFs im ZIP"}


def remove_pages(files, p):
    f = files[0]
    doc = open_pdf(f)
    pages = parse_pages(p.get("pages"), doc.page_count, default_all=False)
    if len(pages) >= doc.page_count:
        raise ToolError("Es können nicht alle Seiten entfernt werden.")
    doc.delete_pages(pages)
    return pdf_result(doc, f"{stem(f['name'])}_bereinigt.pdf",
                      f"{len(pages)} Seite(n) entfernt, {doc.page_count} verbleiben")


def extract_pages(files, p):
    f = files[0]
    doc = open_pdf(f)
    pages = parse_pages(p.get("pages"), doc.page_count, default_all=False)
    if p.get("separate"):
        parts = []
        for i in pages:
            d = pymupdf.open()
            d.insert_pdf(doc, from_page=i, to_page=i)
            parts.append((f"{stem(f['name'])}_S{i + 1}.pdf", save_bytes(d)))
        return {"data": zip_files(parts), "name": f"{stem(f['name'])}_seiten.zip",
                "info": f"{len(parts)} einzelne PDFs"}
    out = pymupdf.open()
    for i in pages:
        out.insert_pdf(doc, from_page=i, to_page=i)
    return pdf_result(out, f"{stem(f['name'])}_auszug.pdf", f"{len(pages)} Seite(n) extrahiert")


def organize(files, p):
    """p['order'] = [{file: idx, page: idx, rotate: deg} | {blank: true}]"""
    docs = [open_pdf(f) for f in files]
    order = p.get("order") or []
    if not order:
        raise ToolError("Keine Seiten übrig.")
    out = pymupdf.open()
    last_size = PAPER["a4"]
    for item in order:
        if item.get("blank"):
            out.new_page(width=last_size[0], height=last_size[1])
            continue
        d = docs[int(item.get("file", 0))]
        i = int(item["page"])
        out.insert_pdf(d, from_page=i, to_page=i)
        pg = out[-1]
        rot = int(item.get("rotate", 0)) % 360
        if rot:
            pg.set_rotation((pg.rotation + rot) % 360)
        last_size = (pg.rect.width, pg.rect.height)
    return pdf_result(out, f"{stem(files[0]['name'])}_organisiert.pdf", f"{out.page_count} Seiten")


def rotate(files, p):
    f = files[0]
    doc = open_pdf(f)
    angle = int(p.get("angle", 90))
    pages = parse_pages(p.get("pages"), doc.page_count)
    for i in pages:
        pg = doc[i]
        pg.set_rotation((pg.rotation + angle) % 360)
    return pdf_result(doc, f"{stem(f['name'])}_gedreht.pdf", f"{len(pages)} Seite(n) um {angle}° gedreht")


COMPRESS = {
    "niedrig": dict(dpi_threshold=200, dpi_target=150, quality=80),
    "empfohlen": dict(dpi_threshold=150, dpi_target=110, quality=65),
    "extrem": dict(dpi_threshold=96, dpi_target=72, quality=45),
}
# Steps for “target size”: downsample images more and more, finally rasterise pages
TARGET_STEPS = [(170, 80), (150, 70), (130, 62), (110, 55), (96, 48), (85, 42), (72, 38), (60, 32), (50, 28)]
RASTER_STEPS = [(110, 55), (96, 48), (80, 42), (68, 36), (56, 30)]


def _optimize(path, gray, dpi_threshold, dpi_target, quality):
    doc = pymupdf.open(path)
    try:
        doc.rewrite_images(dpi_threshold=dpi_threshold, dpi_target=dpi_target, quality=quality,
                           lossy=True, lossless=True, bitonal=True, color=True, gray=True,
                           set_to_gray=bool(gray))
    except Exception:
        pass
    for fn in (doc.subset_fonts, doc.del_xml_metadata):
        try:
            fn()
        except Exception:
            pass
    return doc.tobytes(garbage=4, deflate=True, deflate_images=True, deflate_fonts=True,
                       clean=True, use_objstms=1)


def _rasterize(path, gray, dpi, quality):
    doc = pymupdf.open(path)
    out = pymupdf.open()
    for page in doc:
        pix = page.get_pixmap(dpi=dpi, alpha=False, colorspace=pymupdf.csGRAY if gray else pymupdf.csRGB)
        img = Image.frombytes("L" if gray else "RGB", (pix.width, pix.height), pix.samples)
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=quality, optimize=True)
        np_ = out.new_page(width=page.rect.width, height=page.rect.height)
        np_.insert_image(np_.rect, stream=buf.getvalue())
    return out.tobytes(garbage=4, deflate=True)


def compress(files, p):
    f = files[0]
    level = p.get("level", "empfohlen")
    gray = p.get("gray")
    before = os.path.getsize(f["path"])
    open_pdf(f).close()
    note = ""
    if level == "ziel":
        target = float(str(p.get("target_mb") or 5).replace(",", ".")) * 1024 * 1024
        best = None
        for dpi, q in TARGET_STEPS:
            data = _optimize(f["path"], gray, int(dpi * 1.15), dpi, q)
            if best is None or len(data) < len(best):
                best = data
            if len(data) <= target:
                note = f" · Bilder auf {dpi} dpi"
                break
        else:
            if p.get("allow_raster", True):
                for dpi, q in RASTER_STEPS:
                    data = _rasterize(f["path"], gray, dpi, q)
                    if len(data) < len(best):
                        best = data
                    if len(data) <= target:
                        note = f" · Seiten als Bild mit {dpi} dpi"
                        break
        data = best
        if len(data) > target:
            note += f" · Ziel {human(target)} nicht ganz erreicht – kleiner geht es ohne unlesbare Qualität nicht"
    elif level == "raster":
        data = _rasterize(f["path"], gray, int(p.get("dpi") or 100), 55)
    else:
        data = _optimize(f["path"], gray, **COMPRESS.get(level, COMPRESS["empfohlen"]))
    if len(data) >= before:
        with open(f["path"], "rb") as fh:
            data = fh.read()
        info = "Die Datei ist bereits optimal komprimiert – keine weitere Verkleinerung möglich."
    else:
        info = f"{human(before)} → {human(len(data))} ({100 - len(data) * 100 // max(before, 1)} % kleiner){note}"
    return {"data": data, "name": f"{stem(f['name'])}_komprimiert.pdf", "info": info}


def human(n):
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024


def scan_effect(files, p):
    f = files[0]
    doc = open_pdf(f)
    out = scanify_pdf(doc, p.get("intensity", "mittel"), p.get("color", "gray"), int(p.get("dpi") or 150))
    return pdf_result(out, f"{stem(f['name'])}_scan.pdf", f"{out.page_count} Seite(n) „gescannt“")


def grayscale(files, p):
    f = files[0]
    doc = open_pdf(f)
    dpi = int(p.get("dpi") or 150)
    out = pymupdf.open()
    for page in doc:
        pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY, alpha=False)
        np_ = out.new_page(width=page.rect.width, height=page.rect.height)
        np_.insert_image(np_.rect, pixmap=pix)
    return pdf_result(out, f"{stem(f['name'])}_graustufen.pdf")


def images_to_pdf(files, p):
    out = pymupdf.open()
    for f in files:
        image_to_pdf_page(out, load_image_as_pdf_path(f["path"]), p.get("size", "a4"),
                          p.get("orientation", "auto"), float(p.get("margin") or 0) * MM)
    return pdf_result(out, f"{stem(files[0]['name'])}.pdf", f"{len(files)} Bild(er) → PDF")


def pdf_to_images(files, p):
    f = files[0]
    doc = open_pdf(f)
    fmt = p.get("format", "jpg")
    dpi = int(p.get("dpi") or 150)
    pages = parse_pages(p.get("pages"), doc.page_count)
    imgs = []
    for i in pages:
        pix = doc[i].get_pixmap(dpi=dpi, alpha=False)
        data = pix.tobytes("jpg", jpg_quality=90) if fmt == "jpg" else pix.tobytes("png")
        imgs.append((f"{stem(f['name'])}_S{i + 1}.{fmt}", data))
    if len(imgs) == 1:
        return {"data": imgs[0][1], "name": imgs[0][0]}
    return {"data": zip_files(imgs), "name": f"{stem(f['name'])}_bilder.zip", "info": f"{len(imgs)} Bilder"}


def extract_images(files, p):
    f = files[0]
    doc = open_pdf(f)
    seen, imgs = set(), []
    for pno, page in enumerate(doc):
        for img in page.get_images(full=True):
            xref = img[0]
            if xref in seen:
                continue
            seen.add(xref)
            info = doc.extract_image(xref)
            if not info or info["width"] < 16:
                continue
            imgs.append((f"S{pno + 1}_bild{len(imgs) + 1}.{info['ext']}", info["image"]))
    if not imgs:
        raise ToolError("Im PDF wurden keine eingebetteten Bilder gefunden.")
    return {"data": zip_files(imgs), "name": f"{stem(f['name'])}_bilder.zip", "info": f"{len(imgs)} Bilder"}


def page_numbers(files, p):
    f = files[0]
    doc = open_pdf(f)
    pages = parse_pages(p.get("pages"), doc.page_count)
    pos = p.get("position", "bottom-center")
    size = float(p.get("size") or 10)
    start = int(p.get("start") or 1)
    fmt = p.get("format") or "{n}"
    margin = float(p.get("margin") or 10) * MM
    color = hex_to_rgb(p.get("color"))
    total = len(pages) + start - 1
    for k, i in enumerate(pages):
        page = doc[i]
        page.remove_rotation()
        text = fmt.replace("{n}", str(start + k)).replace("{total}", str(total))
        fargs, fobj = font_for("helv")
        w = fobj.text_length(text, fontsize=size)
        r = page.rect
        vert, horiz = pos.split("-")
        x = {"left": margin, "center": (r.width - w) / 2, "right": r.width - margin - w}[horiz]
        y = margin + size if vert == "top" else r.height - margin
        page.insert_text((x, y), text, fontsize=size, color=color, **fargs)
    return pdf_result(doc, f"{stem(f['name'])}_nummeriert.pdf")


def watermark(files, p, image_bytes=None):
    f = files[0]
    doc = open_pdf(f)
    pages = parse_pages(p.get("pages"), doc.page_count)
    text = p.get("text") or "VERTRAULICH"
    size = float(p.get("size") or 60)
    opacity = float(p.get("opacity") or 30) / 100
    color = hex_to_rgb(p.get("color"), (0.8, 0, 0))
    angle = float(p.get("angle") or 45)
    layout = p.get("layout", "center")
    behind = p.get("layer") == "behind"
    img = None
    if p.get("image"):
        img = base64.b64decode(p["image"].split(",", 1)[-1])
    for i in pages:
        page = doc[i]
        page.remove_rotation()
        r = page.rect
        if layout == "tiled":
            step_x, step_y = max(size * 6, 150), max(size * 3.5, 110)
            centers = [(x, y) for y in frange(step_y / 2, r.height, step_y)
                       for x in frange(step_x / 2, r.width, step_x)]
        else:
            centers = [(r.width / 2, r.height / 2)]
        for cx, cy in centers:
            if img:
                iw = r.width * float(p.get("imgscale") or 40) / 100
                pil = Image.open(io.BytesIO(img))
                ih = iw * pil.height / pil.width
                if opacity < 1:
                    pil = pil.convert("RGBA")
                    a = pil.getchannel("A").point(lambda v: int(v * opacity))
                    pil.putalpha(a)
                    buf = io.BytesIO()
                    pil.save(buf, "PNG")
                    data = buf.getvalue()
                else:
                    data = img
                page.insert_image(pymupdf.Rect(cx - iw / 2, cy - ih / 2, cx + iw / 2, cy + ih / 2),
                                  stream=data, overlay=not behind)
            else:
                fargs, fobj = font_for("hebo")
                tw = fobj.text_length(text, fontsize=size)
                m = pymupdf.Matrix(angle)  # positive = from bottom left to top right
                origin = pymupdf.Point(cx - tw / 2, cy + size * 0.35)
                page.insert_text(origin, text, fontsize=size, color=color, **fargs,
                                 fill_opacity=opacity, stroke_opacity=opacity, overlay=not behind,
                                 morph=(pymupdf.Point(cx, cy), m))
    return pdf_result(doc, f"{stem(f['name'])}_wasserzeichen.pdf")


def frange(a, b, s):
    while a < b:
        yield a
        a += s


def crop(files, p):
    f = files[0]
    doc = open_pdf(f)
    pages = parse_pages(p.get("pages"), doc.page_count)
    t, r_, b, l = (float(p.get(k) or 0) * MM for k in ("top", "right", "bottom", "left"))
    for i in pages:
        page = doc[i]
        page.remove_rotation()
        cb = page.cropbox
        new = pymupdf.Rect(cb.x0 + l, cb.y0 + t, cb.x1 - r_, cb.y1 - b)
        if new.is_empty or new.width < 20 or new.height < 20:
            raise ToolError("Die Ränder sind zu gross für die Seitengrösse.")
        page.set_cropbox(new)
    return pdf_result(doc, f"{stem(f['name'])}_zugeschnitten.pdf")


def protect(files, p):
    f = files[0]
    doc = open_pdf(f)
    user = p.get("password") or ""
    if not user:
        raise ToolError("Bitte ein Passwort eingeben.")
    if user != (p.get("password2") or ""):
        raise ToolError("Die Passwörter stimmen nicht überein.")
    perm = pymupdf.PDF_PERM_ACCESSIBILITY
    if p.get("allow_print", True):
        perm |= pymupdf.PDF_PERM_PRINT | pymupdf.PDF_PERM_PRINT_HQ
    if p.get("allow_copy", True):
        perm |= pymupdf.PDF_PERM_COPY
    if p.get("allow_edit", False):
        perm |= pymupdf.PDF_PERM_MODIFY | pymupdf.PDF_PERM_ANNOTATE | pymupdf.PDF_PERM_FORM | pymupdf.PDF_PERM_ASSEMBLE
    owner = p.get("owner") or (user + "#owner")
    data = doc.tobytes(garbage=3, deflate=True, encryption=pymupdf.PDF_ENCRYPT_AES_256,
                       user_pw=user, owner_pw=owner, permissions=perm)
    return {"data": data, "name": f"{stem(f['name'])}_geschuetzt.pdf", "info": "AES-256 verschlüsselt"}


def unlock(files, p):
    f = files[0]
    doc = pymupdf.open(f["path"])
    if doc.is_encrypted or doc.needs_pass:
        if doc.needs_pass and not doc.authenticate(p.get("password") or ""):
            raise ToolError("Falsches Passwort.")
    data = doc.tobytes(garbage=3, deflate=True, encryption=pymupdf.PDF_ENCRYPT_NONE)
    return {"data": data, "name": f"{stem(f['name'])}_entsperrt.pdf", "info": "Schutz entfernt"}


def redact(files, p):
    f = files[0]
    doc = open_pdf(f)
    terms = [t.strip() for t in re.split(r"[\n,;]", p.get("terms") or "") if t.strip()]
    patterns = []
    if p.get("emails"):
        patterns.append(r"[\w.+-]+@[\w-]+\.[\w.-]+")
    if p.get("iban"):
        patterns.append(r"\b[A-Z]{2}\d{2}(?: ?[A-Z0-9]{4}){2,7}(?: ?[A-Z0-9]{1,4})?\b")
    if p.get("phones"):
        patterns.append(r"\+?\d[\d /-]{7,}\d")
    if not terms and not patterns:
        raise ToolError("Bitte Begriffe oder Muster zum Schwärzen angeben.")
    count = 0
    for page in doc:
        rects = []
        for t in terms:
            rects += page.search_for(t)
        if patterns:
            text = page.get_text()
            for pat in patterns:
                for m in set(re.findall(pat, text)):
                    rects += page.search_for(m.strip())
        for r in rects:
            page.add_redact_annot(r, fill=(0, 0, 0))
        if rects:
            page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_PIXELS)
            count += len(rects)
    if not count:
        raise ToolError("Keine Treffer gefunden – nichts geschwärzt.")
    doc.scrub(metadata=True)
    return pdf_result(doc, f"{stem(f['name'])}_geschwaerzt.pdf", f"{count} Stelle(n) dauerhaft geschwärzt")


def tessdata():
    try:
        return pymupdf.get_tessdata()
    except Exception:
        for d in ("/opt/homebrew/share/tessdata", "/usr/local/share/tessdata",
                  "/usr/share/tesseract-ocr/5/tessdata", "/usr/share/tesseract-ocr/4.00/tessdata"):
            if os.path.isdir(d):
                return d
    return None


def ocr(files, p):
    f = files[0]
    doc = open_pdf(f)
    td = tessdata()
    if not td:
        raise ToolError("Tesseract ist nicht installiert. Im Terminal: brew install tesseract tesseract-lang")
    lang = p.get("lang") or "deu+eng"
    only_empty = p.get("only_empty", True)
    done = 0
    font = pymupdf.Font("helv")
    for page in doc:
        if only_empty and page.get_text().strip():
            continue
        page.remove_rotation()
        tp = page.get_textpage_ocr(dpi=300, full=True, language=lang, tessdata=td)
        words = page.get_text("words", textpage=tp)
        tw = pymupdf.TextWriter(page.rect)
        for w in words:
            x0, y0, x1, y1, word = w[:5]
            h = y1 - y0
            if h <= 0 or not word.strip():
                continue
            size = h * 0.9
            natural = font.text_length(word, fontsize=size) or 1
            size = size * min(2.0, (x1 - x0) / natural)
            try:
                tw.append((x0, y1 - h * 0.2), word, font=font, fontsize=size)
            except Exception:
                pass
        tw.write_text(page, render_mode=3)  # unsichtbar, aber durchsuchbar
        done += 1
    return pdf_result(doc, f"{stem(f['name'])}_ocr.pdf", f"Texterkennung auf {done} Seite(n)")


def repair(files, p):
    f = files[0]
    doc = open_pdf(f)
    data = doc.tobytes(garbage=4, deflate=True, clean=True)
    return {"data": data, "name": f"{stem(f['name'])}_repariert.pdf",
            "info": "PDF neu aufgebaut" + (" (Fehler wurden behoben)" if doc.is_repaired else "")}


def flatten(files, p):
    f = files[0]
    doc = open_pdf(f)
    doc.bake(annots=True, widgets=True)
    return pdf_result(doc, f"{stem(f['name'])}_flach.pdf", "Formulare & Kommentare fest eingebrannt")


WIDGET_KIND = {pymupdf.PDF_WIDGET_TYPE_TEXT: "text", pymupdf.PDF_WIDGET_TYPE_CHECKBOX: "check",
               pymupdf.PDF_WIDGET_TYPE_RADIOBUTTON: "radio", pymupdf.PDF_WIDGET_TYPE_COMBOBOX: "choice",
               pymupdf.PDF_WIDGET_TYPE_LISTBOX: "choice"}


def form_fields(files, p):
    """All fillable fields with their position (as fraction of the page) – for the form view."""
    f = files[0]
    doc = open_pdf(f)
    fields = []
    for pno, page in enumerate(doc):
        W, H = page.rect.width, page.rect.height
        for w in page.widgets():
            kind = WIDGET_KIND.get(w.field_type)
            if not kind:
                continue
            r = w.rect
            item = {"id": str(w.xref), "page": pno, "kind": kind, "name": w.field_name or "",
                    "label": w.field_label or "", "x": r.x0 / W, "y": r.y0 / H, "w": r.width / W, "h": r.height / H,
                    "readonly": bool(w.field_flags & pymupdf.PDF_FIELD_IS_READ_ONLY), "size": w.text_fontsize or 0}
            if kind == "text":
                item["value"] = w.field_value or ""
                item["multiline"] = bool(w.field_flags & pymupdf.PDF_TX_FIELD_IS_MULTILINE)
                item["maxlen"] = w.text_maxlen or 0
            elif kind in ("check", "radio"):
                on = w.on_state()
                item["value"] = bool(w.field_value) and w.field_value not in ("Off", False) and (kind == "check" or w.field_value == on)
            else:
                item["options"] = [o if isinstance(o, str) else o[-1] for o in (w.choice_values or [])]
                item["value"] = w.field_value or ""
            fields.append(item)
    return {"json": {"fields": fields, "pages": doc.page_count}}


def fill_form(files, p):
    f = files[0]
    doc = open_pdf(f)
    values = p.get("values") or {}
    n = 0
    for page in doc:
        widgets = [w for w in page.widgets() if str(w.xref) in values]
        # radio buttons: switch the group off first, then the chosen one on
        for w in widgets:
            if w.field_type == pymupdf.PDF_WIDGET_TYPE_RADIOBUTTON and not values[str(w.xref)]:
                w.field_value = False
                w.update()
        for w in widgets:
            v = values[str(w.xref)]
            if w.field_type == pymupdf.PDF_WIDGET_TYPE_RADIOBUTTON:
                if not v:
                    continue
                w.field_value = w.on_state()
            elif w.field_type == pymupdf.PDF_WIDGET_TYPE_CHECKBOX:
                w.field_value = w.on_state() if v else "Off"
            else:
                w.field_value = str(v)
            w.update()
            n += 1
    if p.get("flatten"):
        doc.bake(annots=False, widgets=True)
    info = f"{n} Felder ausgefüllt" + (" und fest eingebrannt" if p.get("flatten") else "")
    return pdf_result(doc, f"{stem(f['name'])}_ausgefuellt.pdf", info)


def metadata(files, p):
    f = files[0]
    doc = open_pdf(f)
    if p.get("clear"):
        doc.set_metadata({})
        doc.del_xml_metadata()
        info = "Alle Metadaten entfernt"
    else:
        md = doc.metadata or {}
        for k in ("title", "author", "subject", "keywords", "creator"):
            if p.get(k) is not None:
                md[k] = p.get(k)
        md["modDate"] = pymupdf.get_pdf_now()
        doc.set_metadata(md)
        info = "Metadaten aktualisiert"
    return pdf_result(doc, f"{stem(f['name'])}.pdf", info)


def extract_text(files, p):
    f = files[0]
    doc = open_pdf(f)
    parts = []
    for i, page in enumerate(doc):
        parts.append(f"===== Seite {i + 1} =====\n{page.get_text(sort=True)}")
    return {"data": "\n".join(parts).encode("utf-8"), "name": f"{stem(f['name'])}.txt"}


def compare(files, p):
    if len(files) != 2:
        raise ToolError("Bitte genau zwei PDFs zum Vergleichen auswählen.")
    texts = []
    for f in files:
        d = open_pdf(f)
        lines = []
        for i, page in enumerate(d):
            for line in page.get_text(sort=True).splitlines():
                if line.strip():
                    lines.append((i + 1, line.strip()))
        texts.append(lines)
    a, b = texts
    sm = difflib.SequenceMatcher(None, [l for _, l in a], [l for _, l in b], autojunk=False)
    ops, changes = [], 0
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        ops.append({"tag": tag, "a": a[i1:i2], "b": b[j1:j2]})
        if tag != "equal":
            changes += 1
    return {"json": {"ops": ops, "changes": changes, "ratio": round(sm.ratio() * 100, 1),
                     "names": [files[0]["name"], files[1]["name"]]}}


def soffice_path():
    for c in ("/Applications/LibreOffice.app/Contents/MacOS/soffice", shutil.which("soffice") or "",
              shutil.which("libreoffice") or ""):
        if c and os.path.exists(c):
            return c
    return None


def office_to_pdf(files, p):
    exe = soffice_path()
    if not exe:
        raise ToolError("LibreOffice ist nicht installiert (kostenlos: libreoffice.org oder "
                        "brew install --cask libreoffice).")
    outs = []
    for f in files:
        tmp = tempfile.mkdtemp()
        src = os.path.join(tmp, os.path.basename(f["name"]))
        shutil.copy(f["path"], src)
        r = subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
                           capture_output=True, timeout=300)
        pdf = os.path.join(tmp, stem(f["name"]) + ".pdf")
        if not os.path.exists(pdf):
            raise ToolError(f"Konvertierung von „{f['name']}“ fehlgeschlagen: {r.stderr.decode(errors='ignore')[:300]}")
        with open(pdf, "rb") as fh:
            outs.append((stem(f["name"]) + ".pdf", fh.read()))
        shutil.rmtree(tmp, ignore_errors=True)
    if len(outs) == 1:
        return {"data": outs[0][1], "name": outs[0][0]}
    if p.get("merge"):
        out = pymupdf.open()
        for _, data in outs:
            out.insert_pdf(pymupdf.open("pdf", data))
        return pdf_result(out, "konvertiert.pdf", f"{len(outs)} Dateien in ein PDF")
    return {"data": zip_files(outs), "name": "konvertiert.zip", "info": f"{len(outs)} PDFs"}


def pdf_to_word(files, p):
    try:
        from pdf2docx import Converter
    except ImportError:
        raise ToolError("Das Modul pdf2docx fehlt. Im Ordner ausführen: .venv/bin/pip install pdf2docx")
    f = files[0]
    open_pdf(f).close()
    tmp = tempfile.mkdtemp()
    out = os.path.join(tmp, stem(f["name"]) + ".docx")
    cv = Converter(f["path"])
    try:
        cv.convert(out)
    finally:
        cv.close()
    with open(out, "rb") as fh:
        data = fh.read()
    shutil.rmtree(tmp, ignore_errors=True)
    return {"data": data, "name": stem(f["name"]) + ".docx"}


def pdf_to_office(files, p):
    """PDF → PowerPoint/Excel/ODT etc. via LibreOffice (limited layout fidelity)."""
    exe = soffice_path()
    if not exe:
        raise ToolError("Dafür wird LibreOffice benötigt (brew install --cask libreoffice).")
    fmt = p.get("target", "pptx")
    f = files[0]
    if fmt == "pptx":
        # every page as a slide (image) – looks exactly like the PDF
        return _pdf_to_pptx(f)
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "in.pdf")
    shutil.copy(f["path"], src)
    filt = {"xlsx": "xlsx:Calc MS Excel 2007 XML", "odt": "odt"}[fmt]
    infilter = ["--infilter=calc_pdf_Import"] if fmt == "xlsx" else []
    subprocess.run([exe, "--headless", *infilter, "--convert-to", filt, "--outdir", tmp, src],
                   capture_output=True, timeout=300)
    out = os.path.join(tmp, "in." + fmt)
    if not os.path.exists(out):
        raise ToolError("Konvertierung fehlgeschlagen.")
    with open(out, "rb") as fh:
        data = fh.read()
    shutil.rmtree(tmp, ignore_errors=True)
    return {"data": data, "name": f"{stem(f['name'])}.{fmt}"}


def _pdf_to_pptx(f):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        raise ToolError("Das Modul python-pptx fehlt (wird mit pdf2docx normalerweise nicht installiert). "
                        "Im Ordner ausführen: .venv/bin/pip install python-pptx")
    doc = open_pdf(f)
    prs = Presentation()
    r = doc[0].rect
    prs.slide_width = Emu(int(r.width * 12700))
    prs.slide_height = Emu(int(r.height * 12700))
    for page in doc:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        pix = page.get_pixmap(dpi=150, alpha=False)
        slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), 0, 0, prs.slide_width, prs.slide_height)
    buf = io.BytesIO()
    prs.save(buf)
    return {"data": buf.getvalue(), "name": f"{stem(f['name'])}.pptx"}


# ---------------------------------------------------------------- Editor

BROWSER_METRICS = {"he": (0.905, 0.212), "ti": (0.891, 0.216), "co": (0.833, 0.300)}

# Unicode-capable fonts (€, –, „“ …): first the macOS system fonts the editor shows in the
# browser; Liberation (metric-compatible) as Linux/web fallback.
_MAC = "/System/Library/Fonts/Supplemental/"
_LIB = "/usr/share/fonts/truetype/liberation/"
_STYLE = {"lv": ("", "Regular"), "ro": ("", "Regular"), "ur": ("", "Regular"),
          "bo": (" Bold", "Bold"), "it": (" Italic", "Italic"), "bi": (" Bold Italic", "BoldItalic")}
_FAMILY = {"he": ("Arial", "LiberationSans"), "ti": ("Times New Roman", "LiberationSerif"),
           "co": ("Courier New", "LiberationMono")}
_font_cache = {}


def font_for(code):
    """Base14 code (helv, tibo, …) -> (insert_text kwargs, Font object for widths)."""
    if code in _font_cache:
        return _font_cache[code]
    fam_mac, fam_lib = _FAMILY.get(code[:2], _FAMILY["he"])
    mac_sfx, lib_sfx = _STYLE.get(code[2:], ("", "Regular"))
    for path in (f"{_MAC}{fam_mac}{mac_sfx}.ttf", f"/Library/Fonts/{fam_mac}{mac_sfx}.ttf",
                 f"{_LIB}{fam_lib}-{lib_sfx}.ttf", f"/fonts/{fam_lib}-{lib_sfx}.ttf"):
        if os.path.exists(path):
            res = ({"fontname": "W" + code, "fontfile": path}, pymupdf.Font(fontfile=path))
            break
    else:
        res = ({"fontname": code}, pymupdf.Font(code))
    _font_cache[code] = res
    return res


def _font_for(fontname, bold, italic):
    base = {"helv": "he", "tiro": "ti", "cour": "co"}.get(fontname, "he")
    suffix = {(False, False): "lv" if base == "he" else ("ro" if base == "ti" else "ur"),
              (True, False): "bo", (False, True): "it", (True, True): "bi"}[(bool(bold), bool(italic))]
    return base + suffix


def apply_edits(files, p):
    """Insert elements (text, images/signatures, shapes, redactions)."""
    f = files[0]
    doc = open_pdf(f)
    elements = p.get("elements") or []
    touched = {int(e["page"]) for e in elements}
    for i in touched:
        if 0 <= i < doc.page_count:
            doc[i].remove_rotation()
    redacted = set()
    for e in elements:
        i = int(e["page"])
        if not 0 <= i < doc.page_count:
            continue
        page = doc[i]
        W, H = page.rect.width, page.rect.height
        x, y = float(e["x"]) * W, float(e["y"]) * H
        w, h = float(e["w"]) * W, float(e["h"]) * H
        rect = pymupdf.Rect(x, y, x + w, y + h)
        t = e["type"]
        color = hex_to_rgb(e.get("color"), (0, 0, 0))
        opacity = float(e.get("opacity", 1))
        if t == "text":
            size = float(e.get("size", 12))
            font = _font_for(e.get("font", "helv"), e.get("bold"), e.get("italic"))
            lines = (e.get("text") or "").split("\n")
            lh = size * float(e.get("lineHeight", 1.2))
            # baseline as in the browser (CSS line-height, Arial / Times New Roman / Courier New)
            asc, desc = BROWSER_METRICS[font[:2]]
            base_off = (lh - (asc + desc) * size) / 2 + asc * size
            align = e.get("align", "left")
            fargs, fobj = font_for(font)
            for k, line in enumerate(lines):
                lw = fobj.text_length(line, fontsize=size)
                if align == "center":
                    lx = x + (w - lw) / 2
                elif align == "right":
                    lx = x + w - lw
                else:
                    lx = x
                page.insert_text((lx, y + k * lh + base_off), line, fontsize=size, color=color,
                                 fill_opacity=opacity, **fargs)
        elif t == "image":
            data = base64.b64decode(e["image"].split(",", 1)[-1])
            if opacity < 1:
                pil = Image.open(io.BytesIO(data)).convert("RGBA")
                pil.putalpha(pil.getchannel("A").point(lambda v: int(v * opacity)))
                buf = io.BytesIO()
                pil.save(buf, "PNG")
                data = buf.getvalue()
            page.insert_image(rect, stream=data, keep_proportion=False)
        elif t == "whiteout":
            page.draw_rect(rect, color=None, fill=(1, 1, 1), width=0)
        elif t == "rect":
            fill = hex_to_rgb(e["fill"]) if e.get("fill") else None
            page.draw_rect(rect, color=color, fill=fill, width=float(e.get("stroke", 1.5)),
                           stroke_opacity=opacity, fill_opacity=opacity)
        elif t == "ellipse":
            page.draw_oval(rect, color=color, width=float(e.get("stroke", 1.5)), stroke_opacity=opacity)
        elif t == "highlight":
            page.draw_rect(rect, color=None, fill=hex_to_rgb(e.get("color"), (1, 0.9, 0)),
                           fill_opacity=float(e.get("opacity", 0.4)), width=0)
        elif t == "line":
            sw = float(e.get("stroke", 1.5))
            page.draw_line((x, y + h / 2), (x + w, y + h / 2), color=color, width=sw, stroke_opacity=opacity)
        elif t == "check":
            pts = [(x + w * 0.1, y + h * 0.55), (x + w * 0.4, y + h * 0.85), (x + w * 0.9, y + h * 0.15)]
            page.draw_polyline(pts, color=color, width=max(1, min(w, h) * 0.12), closePath=False,
                               lineCap=1, lineJoin=1)
        elif t == "cross":
            sw = max(1, min(w, h) * 0.12)
            page.draw_line((x + w * .15, y + h * .15), (x + w * .85, y + h * .85), color=color, width=sw, lineCap=1)
            page.draw_line((x + w * .85, y + h * .15), (x + w * .15, y + h * .85), color=color, width=sw, lineCap=1)
        elif t == "redact":
            page.add_redact_annot(rect, fill=(0, 0, 0))
            redacted.add(i)
    for i in redacted:
        doc[i].apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_PIXELS)

    name = f"{stem(f['name'])}_bearbeitet.pdf"
    scan = p.get("scan") or {}
    if scan.get("enabled"):
        doc = scanify_pdf(doc, scan.get("intensity", "leicht"), scan.get("color", "gray"),
                          int(scan.get("dpi") or 150))
        name = f"{stem(f['name'])}_unterschrieben_scan.pdf"
    return pdf_result(doc, name, f"{len(elements)} Element(e) eingefügt")


TOOLS = {
    "merge": merge, "split": split, "remove": remove_pages, "extract": extract_pages,
    "organize": organize, "rotate": rotate, "compress": compress, "scan": scan_effect,
    "grayscale": grayscale, "images_to_pdf": images_to_pdf, "pdf_to_images": pdf_to_images,
    "extract_images": extract_images, "page_numbers": page_numbers, "watermark": watermark,
    "crop": crop, "protect": protect, "unlock": unlock, "redact": redact, "ocr": ocr,
    "repair": repair, "flatten": flatten, "metadata": metadata, "extract_text": extract_text,
    "compare": compare, "office_to_pdf": office_to_pdf, "pdf_to_word": pdf_to_word,
    "pdf_to_office": pdf_to_office, "edit": apply_edits,
    "form_fields": form_fields, "fill": fill_form,
}


def capabilities():
    caps = {"ocr": bool(shutil.which("tesseract")), "office": bool(soffice_path())}
    try:
        import pdf2docx  # noqa
        caps["pdf2docx"] = True
    except ImportError:
        caps["pdf2docx"] = False
    try:
        import pptx  # noqa
        caps["pptx"] = True
    except ImportError:
        caps["pptx"] = False
    return caps
